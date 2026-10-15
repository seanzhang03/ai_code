from pydantic import BaseModel,Field
from langchain.tools import tool
from docx import Document
from docx.shared import Pt
import time

class WordParams(BaseModel):
    title:str = Field(...,description="文档标题")
    content:str = Field(...,description="文档内容")

@tool(args_schema=WordParams)  #引入的变量名必须完全一致，进行数据校验
def word_tool(title:str,content:str)->str:
    """
    写入word文档
    """
    try:
        #1.创建一个新的文档对象
        doc = Document()
        #2.添加一级标题
        doc.add_heading(title,level=1)
        #3.添加一个普通段落
        p = doc.add_paragraph(content)
        #定义时间戳作为文件名
        file_name = time.strftime("%Y%m%d%H%M%S",time.localtime())
        doc.save(f"{file_name}.docx")
        return "写入成功"
    except Exception as e:
        print(f"出现异常{e}")
        return "出现异常"